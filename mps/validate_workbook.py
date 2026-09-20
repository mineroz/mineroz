#!/usr/bin/env python3
"""
validate_workbook.py - build, recalculate and independently verify the MPS workbook.

What it does
  1. builds four workbooks with build_workbook.py:
       demo    : 74 days of synthetic DEMO data, report date fixed to 10 Feb
       default : the same data, report date left to the workbook's default formula
       blank   : no data at all
       edge    : the demo with partially entered rows (missing grades, missing run
                 hours, a full-day shutdown, an ungraded ore row, an unknown source,
                 a row dated in the prior year) plus a Probe sheet that measures every
                 lst_ drop-down name
  2. recalculates them with LibreOffice (headless) and reads the values back
  3. recomputes the key figures in plain Python from Builder.sample (the synthetic
     rows, in schema field order) without using any workbook formula, and compares
     them with Calc_Daily for 1 Jan, 10 Feb and the last day with data
  4. checks Daily_Report, Monthly_Summary and Chart_Data against Calc_Daily and
     against the independent figures
  5. scans every sheet for error values and audits the built file with
     tools/xl_inspect.py (no external links, no VBA, no dangling references,
     no broken names)
  6. checks that the blank workbook recalculates without a single error value

Usage (from the repository root):
    python3 mps/validate_workbook.py [--keep] [--workdir DIR] [--soffice PATH]

Prints one line per check (check, expected, actual, pass) and exits with code 1
when any check fails. Temporary files go to a fresh temp folder unless --workdir
is given; --keep leaves them in place.
"""
from __future__ import annotations

import argparse
import calendar
import json
import shutil
import subprocess
import sys
import tempfile
import time
import warnings
from datetime import date, datetime, time as dtime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(HERE))
import build_workbook as bw  # noqa: E402

YEAR = 2026
SAMPLE_DAYS = 74
REPORT_DATE = date(YEAR, 2, 10)
CHECK_DATES = [date(YEAR, 1, 1), REPORT_DATE, date(YEAR, 1, 1) + timedelta(days=SAMPLE_DAYS - 1)]
EDGE_DATES = [date(YEAR, 2, d) for d in (5, 6, 7, 8, 9, 10)]


def set_year(year: int):
    """Point every date expectation at another reporting year (--year 2028 exercises the leap-year branch)."""
    global YEAR, REPORT_DATE, CHECK_DATES, EDGE_DATES
    YEAR = year
    REPORT_DATE = date(YEAR, 2, 10)
    CHECK_DATES = [date(YEAR, 1, 1), REPORT_DATE, date(YEAR, 1, 1) + timedelta(days=SAMPLE_DAYS - 1)]
    EDGE_DATES = [date(YEAR, 2, d) for d in (5, 6, 7, 8, 9, 10)]
TOL_SUM = 1e-6      # relative tolerance for additive figures (tonnes, ounces, hours, counts)
TOL_RATIO = 1e-4    # relative tolerance for ratios, grades, rates and pro-rata budgets
TOL_SAME = 1e-9     # two cells of the same recalculated workbook that must agree
ERROR_VALUES = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#NUM!", "#NULL!", "#N/A")
MONTHS = [calendar.month_abbr[m] for m in range(1, 13)]


# ----------------------------------------------------------------------------- value helpers
def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 20000:
        return from_excel(v).date()
    return None


def norm(v):
    """Normalise a cell value: blank strings become None, datetimes become dates, ints become floats."""
    if v is None:
        return None
    if isinstance(v, str):
        return None if v == "" else v
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return v


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def num(v) -> float:
    """Excel N(): blank -> 0."""
    return float(v) if is_num(v) else 0.0


def close(exp, act, rel) -> bool:
    exp, act = norm(exp), norm(act)
    if exp is None or act is None:
        return exp is None and act is None
    if is_num(exp) and is_num(act):
        return abs(act - exp) <= rel * max(abs(exp), 1e-9)
    return exp == act


def div(n, d):
    """Ratio as the workbook defines it: blank unless the denominator is positive."""
    return n / d if d > 0 else None


def fmt(v) -> str:
    v = norm(v)
    if v is None:
        return "blank"
    if isinstance(v, float):
        if abs(v) < 1e15 and v == int(v):
            return str(int(v))
        return f"{v:.10g}"
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, dtime):
        return "0 (zero date)"
    s = str(v)
    return s if len(s) <= 30 else s[:27] + "..."


class Checks:
    def __init__(self):
        self.rows: list[tuple[str, str, str, bool]] = []

    def add(self, name, expected, actual, ok):
        self.rows.append((name, str(expected), str(actual), bool(ok)))

    def num(self, name, expected, actual, rel=TOL_SUM):
        self.add(name, fmt(expected), fmt(actual), close(expected, actual, rel))

    def eq(self, name, expected, actual):
        self.add(name, fmt(expected), fmt(actual), norm(expected) == norm(actual))

    def group(self, name, total, mismatches):
        exp = f"{total} agree"
        act = exp if not mismatches else f"{len(mismatches)} differ: " + "; ".join(mismatches[:2])
        self.add(name, exp, act, not mismatches)

    @property
    def failed(self):
        return [r for r in self.rows if not r[3]]

    def print(self):
        w0 = min(max(len(r[0]) for r in self.rows), 64)
        w1 = min(max(len(r[1]) for r in self.rows), 30)
        print(f"{'check':<{w0}}  {'expected':<{w1}}  {'actual':<{w1}}  pass")
        print("-" * (w0 + 2 * w1 + 10))
        for name, exp, act, ok in self.rows:
            a = act if ok or len(act) <= 80 else act[:77] + "..."
            print(f"{name[:w0]:<{w0}}  {exp[:w1]:<{w1}}  {a:<{w1}}  {'ok' if ok else 'FAIL'}")
        print("-" * (w0 + 2 * w1 + 10))
        print(f"{len(self.rows)} checks, {len(self.failed)} failed")


# ----------------------------------------------------------------------------- sample access
class Sample:
    """Builder.sample rows (list of lists in schema field order) as dicts keyed by field name."""

    def __init__(self, builder, sample: dict | None = None):
        self.b = builder
        self.sample = sample if sample is not None else builder.sample
        s = builder.s
        self.troy = float(s["constants"]["TROY_OZ_G"])
        self.fields = {t["table"]: [f["name"] for f in t["fields"] if not f.get("calc")] for t in s["tables"]}
        for rdef in s["reference_tables"].values():
            self.fields[rdef["table"]] = [c["name"] for c in rdef["columns"]]
        self.src_type = {r[0]: r[1] for r in s["reference_tables"]["Sources"]["rows"]}
        self.dst_type = {r[0]: r[1] for r in s["reference_tables"]["Destinations"]["rows"]}
        self.width_problems = []
        self._cache = {}

    def rows(self, table) -> list[dict]:
        if table in self._cache:
            return self._cache[table]
        names = self.fields[table]
        out = []
        for row in self.sample.get(table, []):
            if len(row) != len(names):
                self.width_problems.append(f"{table}: row has {len(row)} values for {len(names)} fields")
            out.append({n: (row[i] if i < len(row) else None) for i, n in enumerate(names)})
        self._cache[table] = out
        return out

    def is_ore(self, material) -> bool:
        return str(material or "").startswith("Ore")

    def oz(self, r) -> float:
        """Contained ounces of a movement row: zero when tonnes or grade is blank."""
        if r.get("Tonnes_t") is None or r.get("Grade_gpt") is None:
            return 0.0
        return float(r["Tonnes_t"]) * float(r["Grade_gpt"]) / self.troy

    def last_plant_date(self):
        ds = [r["Date"] for r in self.rows("tblPlant") if r["Milled_t"] is not None]
        return max(ds) if ds else None


def in_period(rd: date, d: date, period: str) -> bool:
    if period == "Day":
        return rd == d
    if period == "MTD":
        return rd.year == d.year and rd.month == d.month and rd <= d
    return rd.year == d.year and rd <= d


def total(rows, d, period, fn) -> float:
    return float(sum(fn(r) for r in rows if in_period(r["Date"], d, period)))


# ----------------------------------------------------------------------------- independent expectations
# key -> True when additive (sum tolerance), False when a ratio (ratio tolerance)
KIND = {
    "Ore_Mined_t": True, "Ore_Mined_oz": True, "Ore_Grade_gpt": False, "Waste_t": True, "TMM_t": True, "Strip_Ratio": False,
    "Rehandle_t": True, "Milled_t": True, "Feed_oz": True, "Recovered_oz": True, "Head_Grade_gpt": False, "Recovery_pct": False,
    "Gold_Poured_oz": True, "Throughput_tph": False, "Mill_Availability_pct": False, "Mill_Utilisation_pct": False,
    "Cyanide_kgpt": False, "Hours_Worked": True, "Recordables": True, "LTI": True, "Gold_Shipped_oz": True, "Gold_Sold_oz": True,
    "GIC_oz": True, "Days_Since_LTI": True, "TRIFR_12m": False, "LTIFR_12m": False,
    "Ore_Ungraded_t": True, "Unclassified_t": True, "Tails_Grade_gpt": False, "Mill_Calendar_h": True,
    "GC_Drill_m": True, "Gravity_Share_pct": False, "Exc_Productivity_tph": False, "Trk_Productivity_tph": False,
    "Bud_Milled_t_MTD": False, "Bud_Milled_t_YTD": False, "Bud_Ore_Grade_gpt_MTD": False, "Bud_Recovery_pct_YTD": False,
}


def expected_kpis(S: Sample, d: date) -> dict[str, dict[str, object]]:
    """Recompute the KPI set for one date, per period, straight from the sample rows.
    Blank means not reported: a row whose grade, tails grade or run hours is blank contributes nothing to the
    ratio that needs it (the tonnes are paired with the assay or the hours that carry them)."""
    T = S.troy
    mv, pl, sf, gd = S.rows("tblMovement"), S.rows("tblPlant"), S.rows("tblSafety"), S.rows("tblGold")
    md, fl = S.rows("tblMiningDaily"), S.rows("tblFleet")

    def pit(r):
        return S.src_type.get(r["Source"]) == "Pit"

    def pit_ore(r):
        return pit(r) and S.is_ore(r["Material"])

    def graded(r):
        return r.get("Grade_gpt") is not None

    def milled(r):
        return num(r["Milled_t"])

    def has(r, *cols):
        return all(r[c] is not None for c in cols)

    def feed_oz(r):
        return milled(r) * num(r["Head_Grade_gpt"]) / T if has(r, "Milled_t", "Head_Grade_gpt") else 0.0

    def tails_oz(r):
        return milled(r) * num(r["Tails_Grade_gpt"]) / T if has(r, "Milled_t", "Tails_Grade_gpt") else 0.0

    def rec_oz(r):
        return feed_oz(r) - tails_oz(r) if has(r, "Milled_t", "Head_Grade_gpt", "Tails_Grade_gpt") else 0.0

    def hours(r):
        return num(r["Hours_Employees"]) + num(r["Hours_Contractors"])

    def recordables(r):
        return num(r["Fatality"]) + num(r["LTI"]) + num(r["RWI"]) + num(r["MTI"])

    out: dict[str, dict[str, object]] = {}

    def put(key, period, val):
        out.setdefault(key, {})[period] = val

    for p in ("Day", "MTD", "YTD"):
        ore_t = total(mv, d, p, lambda r: float(r["Tonnes_t"]) if pit_ore(r) else 0.0)
        ore_graded_t = total(mv, d, p, lambda r: float(r["Tonnes_t"]) if pit_ore(r) and graded(r) else 0.0)
        ore_oz = total(mv, d, p, lambda r: S.oz(r) if pit_ore(r) else 0.0)
        waste_t = total(mv, d, p, lambda r: float(r["Tonnes_t"]) if pit(r) and not S.is_ore(r["Material"]) else 0.0)
        rehandle = total(mv, d, p, lambda r: float(r["Tonnes_t"]) if S.src_type.get(r["Source"]) == "Stockpile" else 0.0)
        moved = total(mv, d, p, lambda r: float(r["Tonnes_t"]))
        mil = total(pl, d, p, milled)
        mil_head = total(pl, d, p, lambda r: milled(r) if r["Head_Grade_gpt"] is not None else 0.0)
        mil_tails = total(pl, d, p, lambda r: milled(r) if r["Tails_Grade_gpt"] is not None else 0.0)
        mil_timed = total(pl, d, p, lambda r: milled(r) if r["Mill_Run_h"] is not None else 0.0)
        fz = total(pl, d, p, feed_oz)
        fz_recon = total(pl, d, p, lambda r: feed_oz(r) if r["Tails_Grade_gpt"] is not None else 0.0)
        tz = total(pl, d, p, tails_oz)
        rz = total(pl, d, p, rec_oz)
        poured = total(pl, d, p, lambda r: num(r["Gold_Poured_oz"]))
        run_h = total(pl, d, p, lambda r: num(r["Mill_Run_h"]))
        plan_h = total(pl, d, p, lambda r: num(r["Mill_Planned_Maint_h"]))
        unpl_h = total(pl, d, p, lambda r: num(r["Mill_Unplanned_Down_h"]))
        cal_h = total(pl, d, p, lambda r: 24.0 if any(r[c] is not None for c in ("Mill_Run_h", "Mill_Planned_Maint_h", "Mill_Unplanned_Down_h")) else 0.0)
        cn_kg = total(pl, d, p, lambda r: num(r["Cyanide_kg"]))
        hrs = total(sf, d, p, hours)
        rec = total(sf, d, p, recordables)
        lti = total(sf, d, p, lambda r: num(r["LTI"]))
        shipped = total(gd, d, p, lambda r: num(r["Gold_oz"]) if r["Type"] == "Shipment" else 0.0)
        gravity = total(pl, d, p, lambda r: num(r["Gravity_Gold_oz"]))
        gc_m = total(md, d, p, lambda r: num(r["GC_Drill_m"]))
        exc_h = total(fl, d, p, lambda r: num(r["Operating_h"]) if r["Equipment_Class"] == "Excavator" else 0.0)
        trk_h = total(fl, d, p, lambda r: num(r["Operating_h"]) if r["Equipment_Class"] == "Haul Truck" else 0.0)
        sold = total(gd, d, p, lambda r: num(r["Gold_oz"]) if r["Type"] == "Sale" else 0.0)
        put("Ore_Mined_t", p, ore_t)
        put("Ore_Mined_oz", p, ore_oz)
        put("Ore_Grade_gpt", p, div(ore_oz * T, ore_graded_t))
        put("Ore_Ungraded_t", p, ore_t - ore_graded_t)
        put("Waste_t", p, waste_t)
        put("TMM_t", p, ore_t + waste_t)
        put("Strip_Ratio", p, div(waste_t, ore_t))
        put("Rehandle_t", p, rehandle)
        put("Unclassified_t", p, moved - ore_t - waste_t - rehandle)
        put("Milled_t", p, mil)
        put("Feed_oz", p, fz)
        put("Recovered_oz", p, rz)
        put("Head_Grade_gpt", p, div(fz * T, mil_head))
        put("Tails_Grade_gpt", p, div(tz * T, mil_tails))
        put("Recovery_pct", p, div(rz, fz_recon))
        put("Gold_Poured_oz", p, poured)
        put("Throughput_tph", p, div(mil_timed, run_h))
        put("Mill_Calendar_h", p, cal_h)
        put("Mill_Availability_pct", p, div(cal_h - plan_h - unpl_h, cal_h))
        put("Mill_Utilisation_pct", p, div(run_h, cal_h - plan_h - unpl_h))
        put("Cyanide_kgpt", p, div(cn_kg, mil))
        put("Hours_Worked", p, hrs)
        put("Recordables", p, rec)
        put("LTI", p, lti)
        put("Gold_Shipped_oz", p, shipped)
        put("GC_Drill_m", p, gc_m)
        put("Gravity_Share_pct", p, div(gravity, rz))
        put("Exc_Productivity_tph", p, div(ore_t + waste_t + rehandle, exc_h))
        put("Trk_Productivity_tph", p, div(ore_t + waste_t + rehandle, trk_h))
        put("Gold_Sold_oz", p, sold)
    # point value: gold in circuit on the day
    gic = total(pl, d, "Day", lambda r: num(r["GIC_oz"]))
    put("GIC_oz", "Day", gic if gic > 0 else None)
    # days since last LTI: prior-year date from Config or the latest LTI day on or before d
    prior = S.b.sample.get("prior_lti")
    lti_days = [r["Date"] for r in sf if num(r["LTI"]) > 0 and r["Date"] <= d]
    cands = [x for x in (prior, max(lti_days) if lti_days else None) if x]
    put("Days_Since_LTI", "Day", float((d - max(cands)).days) if cands else None)
    # rolling 12 months + month to date: window from the same month one year earlier; history months are used
    # for months in the window before the current month that have no daily hours
    hist = S.rows("tblSafetyHistory")
    ms = date(d.year, d.month, 1)
    win = date(ms.year - 1, ms.month, 1)

    def use_hist(h):
        m = h["Month"]
        return win <= m < ms and not any(hours(r) > 0 for r in sf if (r["Date"].year, r["Date"].month) == (m.year, m.month))

    for key, cur_fn, hist_col in (("TRIFR_12m", recordables, "Recordables"), ("LTIFR_12m", lambda r: num(r["LTI"]), "LTI")):
        n = sum(cur_fn(r) for r in sf if win <= r["Date"] <= d) + sum(num(h[hist_col]) for h in hist if use_hist(h))
        dd = sum(hours(r) for r in sf if win <= r["Date"] <= d) + sum(num(h["Hours_Total"]) for h in hist if use_hist(h))
        put(key, "Day", n / dd * 1_000_000 if dd > 0 else None)
    return out


def expected_budgets(S: Sample, d: date) -> dict[str, object]:
    """Pro-rata budgets from tblBudget: month x day-of-month / days in month; YTD = prior full months + MTD."""
    T = S.troy
    bud = {r["Month"]: r for r in S.rows("tblBudget")}
    dim = calendar.monthrange(d.year, d.month)[1]

    def month_val(m, fn):
        r = bud.get(date(d.year, m, 1))
        return fn(r) if r else 0.0

    def mtd(fn):
        return month_val(d.month, fn) * d.day / dim

    def ytd(fn):
        return sum(month_val(m, fn) for m in range(1, d.month)) + mtd(fn)

    milled = lambda r: num(r["Milled_t"])  # noqa: E731
    ore_t = lambda r: num(r["Ore_Mined_t"])  # noqa: E731
    ore_oz = lambda r: num(r["Ore_Mined_t"]) * num(r["Ore_Grade_gpt"]) / T  # noqa: E731
    feed_oz = lambda r: num(r["Milled_t"]) * num(r["Head_Grade_gpt"]) / T  # noqa: E731
    rec_oz = lambda r: feed_oz(r) * num(r["Recovery_pct"])  # noqa: E731
    return {
        "Bud_Milled_t_MTD": mtd(milled),
        "Bud_Milled_t_YTD": ytd(milled),
        "Bud_Ore_Grade_gpt_MTD": div(mtd(ore_oz) * T, mtd(ore_t)),
        "Bud_Recovery_pct_YTD": div(ytd(rec_oz), ytd(feed_oz)),
    }


def expected_stockpiles(S: Sample, rd: date) -> list[dict]:
    """Balance window per stockpile: from the later of 1 January and its Survey_Date (opening = start of that day) to rd."""
    T = S.troy
    out = []
    for sp in S.rows("tblStockpiles"):
        nm, op_t, op_g = sp["Stockpile"], num(sp["Opening_t"]), num(sp["Opening_gpt"])
        since = max(date(rd.year, 1, 1), as_date(sp["Survey_Date"]) or date(rd.year, 1, 1))
        mv = [r for r in S.rows("tblMovement") if since <= r["Date"] <= rd]
        pl = [r for r in S.rows("tblPlant") if since <= r["Date"] <= rd]
        milled = sum(num(r["Milled_t"]) for r in pl)
        feed_oz = sum(num(r["Milled_t"]) * num(r["Head_Grade_gpt"]) / T for r in pl)
        tip = [r for r in mv if S.dst_type.get(r["Destination"]) == "Plant"]
        tip_t = sum(float(r["Tonnes_t"]) for r in tip)
        tip_oz = sum(S.oz(r) for r in tip)
        feeds = sp["Plant_Feed"] == "Y"
        ins = [r for r in mv if r["Destination"] == nm]
        outs = [r for r in mv if r["Source"] == nm]
        in_t, in_oz = sum(float(r["Tonnes_t"]) for r in ins), sum(S.oz(r) for r in ins)
        out_t, out_oz = sum(float(r["Tonnes_t"]) for r in outs), sum(S.oz(r) for r in outs)
        pf_t = milled - tip_t if feeds else 0.0
        pf_oz = feed_oz - tip_oz if feeds else 0.0
        bal = op_t + in_t - out_t - pf_t
        bal_oz = op_t * op_g / T + in_oz - out_oz - pf_oz
        out.append({"name": nm, "op_t": op_t, "op_g": op_g, "in_t": in_t, "out_t": out_t, "pf_t": pf_t, "bal": bal,
                    "grade": bal_oz * T / bal if bal > 0 else None})
    return out


# ----------------------------------------------------------------------------- workbook readers
class CalcDaily:
    def __init__(self, ws):
        rows = list(ws.iter_rows(values_only=True))
        self.hdr = {h: i for i, h in enumerate(rows[0]) if h}
        self.by_date = {}
        for row in rows[1:]:
            d = as_date(row[0])
            if d:
                self.by_date[d] = row

    def get(self, d: date, hdr: str):
        if hdr not in self.hdr:
            return f"<no column {hdr}>"
        row = self.by_date.get(d)
        return norm(row[self.hdr[hdr]]) if row else f"<no row {d}>"


def config_values(wb) -> dict:
    ws = wb["Config"]
    out = {}
    for r in range(1, 40):
        nm = ws.cell(r, 3).value
        if isinstance(nm, str) and nm.startswith("cfg_"):
            out[nm] = ws.cell(r, 2).value
    return out


def last_data_date(cfg: dict):
    d = as_date(cfg.get("cfg_LastDataDate"))
    return d if d and d.year >= 1901 else None


def find_row(ws, col: int, text: str, start=1):
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.startswith(text):
            return r
    return None


def scan_errors(wb, allowed) -> dict[str, list[str]]:
    found: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v in ERROR_VALUES and not allowed(ws.title, c):
                    found.setdefault(ws.title, []).append(f"{c.coordinate}={v}")
    return found


def chart_gap_ok(cd, src: str, d: date, last) -> bool:
    """A guarded Chart_Data cell may be #N/A after the last data date, or where the Calc_Daily value is not a number (blank ratio)."""
    return last is None or d > last or (cd is not None and not is_num(cd.get(d, src)))


def chart_guard_allowed(wb, series, last, cd=None):
    """Only #N/A in guarded Chart_Data columns is acceptable, and only where chart_gap_ok says so."""
    ws = wb["Chart_Data"]
    guarded = {j + 1: s[1] for j, s in enumerate(series) if s[2]}

    def allowed(sheet, c):
        if sheet != "Chart_Data" or c.value != "#N/A" or c.column not in guarded:
            return False
        d = as_date(ws.cell(c.row, 1).value)
        return d is not None and chart_gap_ok(cd, guarded[c.column], d, last)

    return allowed


# ----------------------------------------------------------------------------- checks
def check_calc_daily(C: Checks, S: Sample, cd: CalcDaily, dates=None, tag="Calc_Daily", budgets=True):
    for d in dates or CHECK_DATES:
        exp = expected_kpis(S, d)
        for key, periods in exp.items():
            for period, val in periods.items():
                hdr = key if period == "Day" else f"{key}_{period}"
                C.num(f"{tag} {d} {hdr}", val, cd.get(d, hdr), TOL_SUM if KIND[key] else TOL_RATIO)
        if budgets:
            for hdr, val in expected_budgets(S, d).items():
                C.num(f"{tag} {d} {hdr}", val, cd.get(d, hdr), TOL_RATIO)


def expected_checks_line(S: Sample, rd: date) -> str:
    """The Daily_Report B7 completeness line, from the sample rows."""
    def n(table, cond=lambda r: True):
        return sum(1 for r in S.rows(table) if r["Date"] == rd and cond(r))
    mv = S.rows("tblMovement")
    unknown = sum(1 for r in mv if r["Source"] and r["Source"] not in S.src_type) + sum(1 for r in mv if r["Destination"] and r["Destination"] not in S.dst_type)
    outside = sum(1 for r in mv if r["Date"].year != rd.year)
    return (f"Rows for this date: Plant {n('tblPlant', lambda r: r['Milled_t'] is not None)}, Movements {n('tblMovement')}, Fleet {n('tblFleet')}, "
            f"Mining daily {n('tblMiningDaily', lambda r: r['Diesel_L'] is not None)}, Safety {n('tblSafety', lambda r: r['Hours_Employees'] is not None)}, "
            f"Comments {n('tblCommentary')}  |  Movement rows with unknown source or destination: {unknown}  |  Movement rows dated outside the year: {outside}")


def var(a, b):
    return a / b - 1 if is_num(a) and is_num(b) and b != 0 else None


def check_daily_report(C: Checks, S: Sample, b, wb, cd: CalcDaily, rd: date, tag="Daily_Report"):
    ws = wb["Daily_Report"]
    C.eq(f"{tag} C5 report date", rd, ws["C5"].value)
    C.eq(f"{tag} rpt_Row (L5) = day of year", float((rd - date(YEAR, 1, 1)).days + 1), ws["L5"].value)
    C.eq(f"{tag} B7 completeness line (rpt_Checks)", expected_checks_line(S, rd), ws["B7"].value)
    C.eq(f"{tag} F5 warning blank (date in year, year check OK)", None, norm(ws["F5"].value))
    C.add(f"{tag} B6 header shows data-to date and week", "Data to ... | Week ...", fmt(ws["B6"].value),
          isinstance(ws["B6"].value, str) and ws["B6"].value.startswith("Data to ") and "Week " in ws["B6"].value)
    C.eq(f"{tag} cfg_FeedCheck (exactly one Plant_Feed = Y)", "OK", config_values(wb).get("cfg_FeedCheck"))
    r0 = find_row(ws, 2, "STOCKPILES")
    C.eq(f"{tag} plant feed warning line blank", None, ws.cell(r0 - 1, 2).value if r0 else "<no STOCKPILES block>")
    rows: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if isinstance(v, str) and v not in rows:
            rows[v] = r
    for section, items in bw.REPORT_LAYOUT:
        mism, n = [], 0
        for key, show_budget in items:
            k = bw.KPI[key]
            r = rows.get(k["label"])
            if r is None:
                mism.append(f"row '{k['label']}' missing")
                continue
            has_p = k["kind"] in ("sum", "ratio")
            has_b = show_budget and b.has_budget(k)
            want = {
                4: cd.get(rd, key),
                5: cd.get(rd, key + "_MTD") if has_p else None,
                8: cd.get(rd, key + "_YTD") if has_p else None,
                6: cd.get(rd, f"Bud_{key}_MTD") if has_b else None,
                9: cd.get(rd, f"Bud_{key}_YTD") if has_b else None,
            }
            want[7] = var(want[5], want[6]) if has_b else None
            want[10] = var(want[8], want[9]) if has_b else None
            for col, w in want.items():
                n += 1
                got = norm(ws.cell(r, col).value)
                if not close(w, got, TOL_SAME):
                    mism.append(f"{k['label']} {get_column_letter(col)}{r} exp {fmt(w)} got {fmt(got)}")
        C.group(f"{tag} {section} cells vs Calc_Daily", n, mism)
    # stockpiles, recomputed independently
    r0 = find_row(ws, 2, "STOCKPILES")
    exp_sp = expected_stockpiles(S, rd)
    mism, n = [], 0
    plaus = []
    for i, e in enumerate(exp_sp):
        r = (r0 or 0) + 1 + i
        got = {c: norm(ws.cell(r, c).value) for c in range(2, 11)} if r0 else {}
        for col, key, tol in ((2, "name", 0), (4, "op_t", TOL_SUM), (5, "op_g", TOL_RATIO), (6, "in_t", TOL_SUM), (7, "out_t", TOL_SUM),
                              (8, "pf_t", TOL_SUM), (9, "bal", TOL_SUM), (10, "grade", TOL_RATIO)):
            n += 1
            if not close(e[key], got.get(col), tol):
                mism.append(f"{e['name']} {get_column_letter(col)}{r} exp {fmt(e[key])} got {fmt(got.get(col))}")
        if not (e["bal"] >= 0 and (e["grade"] is None or 0 <= e["grade"] <= 10)):
            plaus.append(f"{e['name']} balance {fmt(e['bal'])} t at {fmt(e['grade'])} g/t")
    C.group(f"{tag} stockpile block (independent)", n, mism)
    C.group(f"{tag} stockpile balances plausible (t >= 0, 0 to 10 g/t)", len(exp_sp), plaus)
    # commentary in entry order
    r0 = find_row(ws, 2, "COMMENTARY")
    comments = [(c["Area"], c["Comment"]) for c in S.rows("tblCommentary") if c["Date"] == rd]
    mism, n = [], 0
    for i in range(bw.COMMENT_LINES):
        e = comments[i] if i < len(comments) else (None, None)
        r = (r0 or 0) + 1 + i
        got = (norm(ws.cell(r, 2).value), norm(ws.cell(r, 3).value)) if r0 else (None, None)
        n += 1
        if got != e:
            mism.append(f"line {i + 1} exp {e} got {got}")
    C.add(f"{tag} commentary lines for report date", f"{len(comments)} comments, entry order", f"{len(comments)} comments" if not mism else mism[0], not mism)
    C.add(f"{tag} commentary count fits the block", f"<= {bw.COMMENT_LINES}", str(len(comments)), len(comments) <= bw.COMMENT_LINES)


def check_monthly(C: Checks, S: Sample, b, wb, cd: CalcDaily, last: date):
    ws = wb["Monthly_Summary"]
    actual_rows: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        v, kind = ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(v, str) and kind == "Actual" and v not in actual_rows:
            actual_rows[v] = r
    # Milled_t, independent of the workbook
    r = actual_rows[bw.KPI["Milled_t"]["label"]]
    pl = S.rows("tblPlant")
    bud = {row["Month"]: row for row in S.rows("tblBudget")}
    # budget: full month, except the month in progress (pro-rata to the last data date); year = YTD budget at that date
    exp_act, exp_bud = {}, {}
    for m in range(1, 13):
        ms = date(YEAR, m, 1)
        exp_act[m] = sum(num(p["Milled_t"]) for p in pl if p["Date"].month == m and p["Date"] <= last) if ms <= last else None
        full = num(bud[ms]["Milled_t"]) if ms in bud else 0.0
        exp_bud[m] = full * last.day / calendar.monthrange(YEAR, m)[1] if m == last.month else full
    for m in (1, 2, 3, 4):
        C.num(f"Monthly_Summary Milled {MONTHS[m - 1]} actual", exp_act[m], ws.cell(r, 3 + m).value)
    ytd = sum(num(p["Milled_t"]) for p in pl if p["Date"] <= last)
    C.num("Monthly_Summary Milled Year actual = YTD at last data date", ytd, ws.cell(r, 16).value)
    mism_b, mism_v = [], []
    for m in range(1, 13):
        got_b = norm(ws.cell(r + 1, 3 + m).value)
        if not close(exp_bud[m], got_b, TOL_RATIO):
            mism_b.append(f"{MONTHS[m - 1]} exp {fmt(exp_bud[m])} got {fmt(got_b)}")
        got_v = norm(ws.cell(r + 2, 3 + m).value)
        ev = var(exp_act[m], exp_bud[m])
        if not close(ev, got_v, TOL_RATIO):
            mism_v.append(f"{MONTHS[m - 1]} exp {fmt(ev)} got {fmt(got_v)}")
    ytd_bud = sum(v for m, v in exp_bud.items() if m <= last.month)
    if not close(ytd_bud, ws.cell(r + 1, 16).value, TOL_RATIO):
        mism_b.append(f"Year exp {fmt(ytd_bud)} got {fmt(ws.cell(r + 1, 16).value)}")
    if not close(var(ytd, ytd_bud), ws.cell(r + 2, 16).value, TOL_RATIO):
        mism_v.append(f"Year exp {fmt(var(ytd, ytd_bud))} got {fmt(ws.cell(r + 2, 16).value)}")
    C.group("Monthly_Summary Milled budget row = tblBudget, pro-rata for the current month (12 months + year)", 13, mism_b)
    C.group("Monthly_Summary Milled variance row = actual / budget - 1", 13, mism_v)
    # every monthly KPI row against Calc_Daily
    mism, n = [], 0
    for key in bw.MONTHLY_KPIS:
        k = bw.KPI[key]
        r = actual_rows.get(k["label"])
        if r is None:
            mism.append(f"row '{k['label']}' missing")
            continue
        per = k["kind"] in ("sum", "ratio")
        for m in range(1, 13):
            ms = date(YEAR, m, 1)
            me = min(date(YEAR, m, calendar.monthrange(YEAR, m)[1]), last)
            want = cd.get(me, f"{key}_MTD" if per else key) if ms <= last else None
            n += 1
            got = norm(ws.cell(r, 3 + m).value)
            if not close(want, got, TOL_SAME):
                mism.append(f"{k['label']} {MONTHS[m - 1]} exp {fmt(want)} got {fmt(got)}")
            if b.has_budget(k):
                n += 1
                wb_ = cd.get(me, f"Bud_{key}_MTD") if ms <= last else cd.get(ms, f"Bud_{key}_Month")
                got_b = norm(ws.cell(r + 1, 3 + m).value)
                if not close(wb_, got_b, TOL_SAME):
                    mism.append(f"{k['label']} {MONTHS[m - 1]} budget exp {fmt(wb_)} got {fmt(got_b)}")
        n += 1
        want = cd.get(last, f"{key}_YTD" if per else key)
        got = norm(ws.cell(r, 16).value)
        if not close(want, got, TOL_SAME):
            mism.append(f"{k['label']} Year exp {fmt(want)} got {fmt(got)}")
    C.group("Monthly_Summary all KPI rows vs Calc_Daily", n, mism)


def check_chart_data(C: Checks, wb, series, last: date, cd=None):
    ws = wb["Chart_Data"]
    rows = list(ws.iter_rows(values_only=True))
    for j, (name, src, guard) in enumerate(series):
        if j == 0:
            continue
        bad = []
        for row in rows[1:]:
            d = as_date(row[0])
            if d is None:
                continue
            v = row[j]
            if guard and (d > last or (cd is not None and not is_num(cd.get(d, src)))):
                if v != "#N/A":
                    bad.append(f"{d} {fmt(v)}")
            elif not is_num(v):
                bad.append(f"{d} {fmt(v)}")
        label = "NA after last data date or blank KPI, numeric before" if guard else "numeric all year"
        C.group(f"Chart_Data {name} ({label})", len(rows) - 1, bad)


def check_inspect(C: Checks, built: Path, outdir: Path):
    cmd = [sys.executable, str(REPO / "tools" / "xl_inspect.py"), str(built), "-o", str(outdir)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    js = outdir / (built.stem + ".json")
    if res.returncode != 0 or not js.exists():
        C.add("xl_inspect ran", "exit 0 and JSON written", f"exit {res.returncode}: {res.stderr.strip()[-100:]}", False)
        return
    info = json.loads(js.read_text(encoding="utf-8"))
    C.add("xl_inspect external links", "0", str(len(info.get("external_links", []))), not info.get("external_links"))
    C.add("xl_inspect VBA project", "none", "present" if info.get("has_vba_project") else "none", not info.get("has_vba_project"))
    dang = info.get("dangling_sheet_refs", [])
    C.add("xl_inspect dangling sheet references", "0", str(len(dang)) + (f" ({', '.join(dang[:3])})" if dang else ""), not dang)
    broken = [d["name"] for d in info.get("defined_names", []) if d.get("broken")]
    C.add("xl_inspect broken defined names", "0", str(len(broken)) + (f" ({', '.join(broken[:3])})" if broken else ""), not broken)
    C.add("xl_inspect scan errors", "none", "; ".join(info.get("errors", [])) or "none", not info.get("errors"))


def check_blank(C: Checks, wb, series):
    cfg = config_values(wb)
    last = last_data_date(cfg)
    C.add("Blank cfg_LastDataDate", "no data (blank or zero)", fmt(cfg.get("cfg_LastDataDate")), last is None)
    found = scan_errors(wb, chart_guard_allowed(wb, series, last))
    C.add("Blank workbook error values outside the Chart_Data guard", "0",
          str(sum(len(v) for v in found.values())) + ("; ".join(f"{k}: {v[0]}" for k, v in found.items())[:60] if found else ""), not found)
    ws = wb["Daily_Report"]
    title = norm(ws["B2"].value)
    C.add("Blank Daily_Report title renders", "text", fmt(title), isinstance(title, str) and bool(title))
    C.eq("Blank cfg_FeedCheck (schema stockpiles)", "OK", cfg.get("cfg_FeedCheck"))
    C.add("Blank Daily_Report B7 completeness line all zero", "Plant 0 ... outside the year: 0", fmt(ws["B7"].value),
          isinstance(ws["B7"].value, str) and ws["B7"].value.startswith("Rows for this date: Plant 0, Movements 0, Fleet 0, Mining daily 0, Safety 0, Comments 0")
          and ws["B7"].value.endswith("outside the year: 0"))
    today = date.today()
    exp_default = max(date(YEAR, 1, 1), min(today - timedelta(days=1), date(YEAR, 12, 31)))
    c5 = as_date(ws["C5"].value)
    C.eq("Blank Daily_Report default date = clamp(today - 1) in year", exp_default, c5)
    C.eq("Blank Daily_Report rpt_Row resolves", float((exp_default - date(YEAR, 1, 1)).days + 1), ws["L5"].value)
    bad, n = [], 0
    for section, items in bw.REPORT_LAYOUT:
        for key, _sb in items:
            r = find_row(ws, 2, bw.KPI[key]["label"])
            if r is None:
                bad.append(f"row '{bw.KPI[key]['label']}' missing")
                continue
            for col in range(4, 11):
                n += 1
                v = norm(ws.cell(r, col).value)
                if v is not None and v != 0:
                    bad.append(f"{get_column_letter(col)}{r}={fmt(v)}")
    C.group("Blank Daily_Report KPI cells blank or zero", n, bad)


# ----------------------------------------------------------------------------- edge workbook and list probe
def list_catalogue(schema: dict) -> dict[str, list]:
    """Every lst_ name the generator defines: schema lists, derived lists and the key column of Lists-sheet reference tables."""
    out = dict(schema["lists"])
    out.update(schema.get("derived_lists", {}))
    for rdef in schema["reference_tables"].values():
        if rdef["sheet"] == "Lists":
            out[rdef["key"]] = [row[0] for row in rdef["rows"]]
    return out


def make_edge(b, out: Path) -> Sample:
    """Copy the demo builder's workbook and sample with partially entered rows, add the lst_ probe sheet, save."""
    import copy
    smp = copy.deepcopy(b.sample)
    S = Sample(b, smp)
    wb = b.wb
    edits = {  # table -> (row date, {field: value})
        "tblPlant": [(date(YEAR, 2, 5), {"Tails_Grade_gpt": None}), (date(YEAR, 2, 6), {"Head_Grade_gpt": None}),
                     (date(YEAR, 2, 7), {"Mill_Run_h": None}),
                     (date(YEAR, 2, 8), {"Milled_t": 0, "Mill_Run_h": None, "Mill_Planned_Maint_h": 24, "Mill_Unplanned_Down_h": 0})],
    }
    for tbl, items in edits.items():
        info = b.tables[tbl]
        ws = wb[info["sheet"]]
        names = S.fields[tbl]
        for d, vals in items:
            r = info["first"] + (d - date(YEAR, 1, 1)).days
            row = smp[tbl][(d - date(YEAR, 1, 1)).days]
            assert row[0] == d
            for k, v in vals.items():
                row[names.index(k)] = v
                ws[f"{info['cols'][k]}{r}"] = v
    # ungraded ore row on 9 Feb (first Petowal Pit Ore HG row of that day), an unknown source and a prior-year row
    info = b.tables["tblMovement"]
    ws = wb[info["sheet"]]
    names = S.fields["tblMovement"]
    mv = smp["tblMovement"]
    i9 = next(i for i, row in enumerate(mv) if row[0] == date(YEAR, 2, 9) and row[3] == "Ore HG")
    mv[i9][names.index("Grade_gpt")] = None
    ws[f"{info['cols']['Grade_gpt']}{info['first'] + i9}"] = None
    extra = [[date(YEAR - 1, 12, 31), None, "Petowal Pit", "Ore HG", "ROM Pad", 777, 2.0, None, None, None],
             [date(YEAR, 2, 10), None, "Unknown Pit", "Ore HG", "ROM Pad", 1000, 2.0, None, None, None]]
    for row in extra:
        r = info["first"] + len(mv)
        mv.append(row)
        for k, v in zip(names, row):
            if v is not None:
                ws[f"{info['cols'][k]}{r}"] = v
    # re-surveyed stockpile: Stockpile LG opening reset on 1 Feb (movements and plant rows before that day leave its balance)
    info = b.tables["tblStockpiles"]
    ws = wb[info["sheet"]]
    i_lg = next(i for i, row in enumerate(smp["tblStockpiles"]) if row[0] == "Stockpile LG")
    names = S.fields["tblStockpiles"]
    for k, v in (("Opening_t", 160000), ("Opening_gpt", 0.58), ("Survey_Date", date(YEAR, 2, 1))):
        smp["tblStockpiles"][i_lg][names.index(k)] = v
        ws[f"{info['cols'][k]}{info['first'] + i_lg}"] = v
    # probe sheet: size, first and last entry of every drop-down name
    pr = wb.create_sheet("Probe")
    pr.append(["name", "rows", "first", "last"])
    for lname in list_catalogue(b.s):
        pr.append([lname, f"=ROWS(lst_{lname})", f"=INDEX(lst_{lname},1)", f"=INDEX(lst_{lname},ROWS(lst_{lname}))"])
    wb.save(out)
    S._cache.clear()
    return S


def check_lists(C: Checks, b, wb_edge):
    """Every lst_ name covers exactly the schema list (F001), and every list field validates against its name."""
    cat = list_catalogue(b.s)
    got = {row[0]: row[1:] for row in wb_edge["Probe"].iter_rows(min_row=2, values_only=True) if row[0]}
    bad = []
    for lname, values in cat.items():
        g = got.get(lname)
        exp = (float(len(values)), values[0], values[-1])
        if g is None or tuple(norm(x) for x in g) != exp:
            bad.append(f"lst_{lname} exp {exp} got {g}")
    C.group("Lists: ROWS/first/last of every lst_ name = schema list (LibreOffice)", len(cat), bad)
    bad, n = [], 0
    fields = [(t["table"], f) for t in b.s["tables"] for f in t["fields"]]
    fields += [(rdef["table"], c) for rdef in b.s["reference_tables"].values() for c in rdef["columns"]]
    for tbl, f in fields:
        if f.get("type") != "list":
            continue
        n += 1
        info = b.tables[tbl]
        ws = b.wb[info["sheet"]]
        want = f"=lst_{f['list']}"
        col_ref = f"{info['cols'][f['name']]}{info['first']}"
        dvs = [dv for dv in ws.data_validations.dataValidation if dv.formula1 == want and col_ref in dv.sqref]
        if not dvs:
            bad.append(f"{tbl}[{f['name']}] has no list validation {want} on {col_ref}")
        elif f"lst_{f['list']}" not in b.wb.defined_names:
            bad.append(f"{tbl}[{f['name']}]: name lst_{f['list']} missing")
    C.group("Lists: every list field validates against an existing lst_ name", n, bad)
    req = [f"{r}.Type" for r, rdef in b.s["reference_tables"].items() if r in ("Sources", "Destinations")
           and not any(c["name"] == "Type" and c.get("required") for c in rdef["columns"])]
    C.group("Schema: Sources.Type and Destinations.Type are required", 2, [f"{x} not required" for x in req])


# ----------------------------------------------------------------------------- driver
def recalc(files: list[Path], outdir: Path, soffice: str) -> list[Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    profile = (outdir.parent / "lo_profile").resolve()
    cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir)] + [str(f) for f in files]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=900)
    out = [outdir / f.name for f in files]
    missing = [p for p in out if not p.exists()]
    if missing:
        sys.exit(f"LibreOffice did not produce {', '.join(str(m) for m in missing)}\n{res.stdout}\n{res.stderr}")
    return out


def load(path: Path):
    return openpyxl.load_workbook(path, data_only=True)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--keep", action="store_true", help="keep the built and recalculated files")
    ap.add_argument("--workdir", default=None, help="folder for the built files (default: a fresh temp folder)")
    ap.add_argument("--soffice", default=shutil.which("soffice") or shutil.which("libreoffice") or "soffice", help="LibreOffice executable")
    ap.add_argument("--schema", default=str(bw.DEFAULT_SCHEMA))
    ap.add_argument("--year", type=int, default=YEAR, help="reporting year to build and check (2028 exercises the leap-year branch)")
    a = ap.parse_args(argv)
    set_year(a.year)

    t0 = time.time()
    auto = a.workdir is None
    work = Path(a.workdir) if a.workdir else Path(tempfile.mkdtemp(prefix="mps_validate_"))
    work.mkdir(parents=True, exist_ok=True)
    built = work / "built"
    built.mkdir(exist_ok=True)
    schema = Path(a.schema)
    C = Checks()
    try:
        print(f"work folder: {work}")
        print("building workbooks ...", flush=True)
        f_demo, f_default, f_blank, f_edge = built / "MPS_demo.xlsx", built / "MPS_demo_default_date.xlsx", built / "MPS_blank.xlsx", built / "MPS_edge.xlsx"
        b = bw.build(schema, f_demo, YEAR, SAMPLE_DAYS, REPORT_DATE)
        bw.build(schema, f_default, YEAR, SAMPLE_DAYS, None)
        bw.build(schema, f_blank, YEAR, 0, None)
        S = Sample(b)
        S_edge = make_edge(b, f_edge)   # after Sample(b): the edge edits the builder's workbook in place
        last = S.last_plant_date()
        C.eq("Sample last day with data", CHECK_DATES[-1], last)
        # the Config cell must hold the formula, not a pasted date (the earlier override bug)
        ref = b.wb.defined_names["cfg_LastDataDate"].attr_text.split("!")[1].replace("$", "")
        cell = b.wb["Config"][ref].value
        C.add("Config cfg_LastDataDate is a formula", "=...", str(cell)[:30], isinstance(cell, str) and cell.startswith("="))
        dvs = [dv for dv in b.wb["Daily_Report"].data_validations.dataValidation if "C5" in str(dv.sqref)]
        C.add("Daily_Report C5 stop validation: date within the year", "date, stop, cfg_YearStart..cfg_YearEnd",
              f"{dvs[0].type}, {dvs[0].errorStyle}, {dvs[0].formula1}..{dvs[0].formula2}" if dvs else "none",
              bool(dvs) and dvs[0].type == "date" and dvs[0].errorStyle == "stop" and dvs[0].formula1 == "=cfg_YearStart" and dvs[0].formula2 == "=cfg_YearEnd")
        C.add("Daily_Report C5 default formula matches the restore hint", "same text", "same" if dvs and bw.DEFAULT_DATE_FORMULA in (dvs[0].prompt or "") else "differs",
              bool(dvs) and bw.DEFAULT_DATE_FORMULA in (dvs[0].prompt or ""))
        for t in S.fields:
            S.rows(t)
        C.group("Sample rows match schema field order (width)", len(S.fields), S.width_problems)

        print("recalculating with LibreOffice ...", flush=True)
        r_demo, r_default, r_blank, r_edge = recalc([f_demo, f_default, f_blank, f_edge], work / "recalc", a.soffice)

        print("checking the DEMO workbook ...", flush=True)
        wb = load(r_demo)
        cfg = config_values(wb)
        C.eq("Recalculated cfg_LastDataDate = last plant date", last, last_data_date(cfg))
        cd = CalcDaily(wb["Calc_Daily"])
        check_calc_daily(C, S, cd)
        check_daily_report(C, S, b, wb, cd, REPORT_DATE)
        check_monthly(C, S, b, wb, cd, last)
        check_chart_data(C, wb, b.chart_series, last, cd)
        found = scan_errors(wb, chart_guard_allowed(wb, b.chart_series, last, cd))
        C.add("DEMO error values outside the Chart_Data guard", "0",
              str(sum(len(v) for v in found.values())) + (" " + "; ".join(f"{k}: {v[0]}" for k, v in found.items())[:60] if found else ""), not found)
        wb.close()

        print("checking the default report date ...", flush=True)
        wb = load(r_default)
        today = date.today()
        exp_default = max(date(YEAR, 1, 1), min(today - timedelta(days=1), date(YEAR, 12, 31), last))
        C.eq("Default report date = min(today - 1, last data date) in year", exp_default, as_date(wb["Daily_Report"]["C5"].value))
        wb.close()

        print("checking the blank workbook ...", flush=True)
        wb = load(r_blank)
        check_blank(C, wb, b.chart_series)
        wb.close()

        print("checking the edge workbook (partial rows) and the drop-down lists ...", flush=True)
        wb = load(r_edge)
        cd = CalcDaily(wb["Calc_Daily"])
        check_calc_daily(C, S_edge, cd, EDGE_DATES, "Edge Calc_Daily", budgets=False)
        check_daily_report(C, S_edge, b, wb, cd, REPORT_DATE, tag="Edge Daily_Report")
        found = scan_errors(wb, chart_guard_allowed(wb, b.chart_series, last, cd))
        C.add("Edge error values outside the Chart_Data guard", "0",
              str(sum(len(v) for v in found.values())) + (" " + "; ".join(f"{k}: {v[0]}" for k, v in found.items())[:60] if found else ""), not found)
        check_lists(C, b, wb)
        wb.close()

        print("running xl_inspect ...", flush=True)
        check_inspect(C, f_demo, work / "inspect")
    finally:
        print()
        C.print()
        print(f"elapsed {time.time() - t0:.0f} s")
        if auto and not a.keep:
            shutil.rmtree(work, ignore_errors=True)
        else:
            print(f"files kept in {work}")
    return 1 if C.failed else 0


if __name__ == "__main__":
    sys.exit(main())
