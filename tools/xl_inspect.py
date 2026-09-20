#!/usr/bin/env python3
"""
xl_inspect.py - Structural audit of Excel workbooks (.xlsx / .xlsm).

Purpose: inventory everything that makes a workbook fragile before rebuilding it
from scratch: external links, VBA, defined names, hidden sheets, cross-sheet
dependency graph, volatile functions, broken references, data connections.

Runs in streaming mode (openpyxl read-only) so multi-hundred-MB workbooks are
handled in bounded memory. Sheet metadata (validations, conditional formats,
merges, protection, freeze panes, declared dimension) is read directly from the
sheet XML in chunks.

Usage:
    python xl_inspect.py <folder-or-file> [more files...] [-o OUTPUT_DIR]

Outputs (in OUTPUT_DIR, default ./xl_inspect_out):
    inspection_report.md      human-readable report
    <workbook>.json           machine-readable detail per workbook

Dependencies: openpyxl (required), oletools (optional, for VBA code analysis).
"""

from __future__ import annotations

import argparse
import json
import os
import posixpath
import re
import sys
import time
import warnings
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install openpyxl")

try:
    from oletools.olevba import VBA_Parser  # type: ignore

    HAVE_OLEVBA = True
except Exception:  # pragma: no cover
    HAVE_OLEVBA = False


NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}

VOLATILE = ("OFFSET(", "INDIRECT(", "NOW(", "TODAY(", "RAND(", "RANDBETWEEN(", "CELL(", "INFO(")
LOOKUPS = ("VLOOKUP(", "HLOOKUP(", "XLOOKUP(", "INDEX(", "MATCH(", "SUMIFS(", "SUMIF(", "COUNTIFS(", "SUMPRODUCT(", "GETPIVOTDATA(")
ERROR_VALUES = ("#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!")
BLOAT_MIN_ROWS = 10000  # declared rows above this, and > 2x the real used rows, is flagged as bloat

# 'Sheet Name'!A1  or  SheetName!A1 ; skip [n]Sheet!A1 (external) and #REF!
SHEET_REF_RE = re.compile(r"(?<![\]#A-Za-z0-9_\.])(?:'((?:[^']|'')+)'|([A-Za-z0-9_\.]+))!")
EXT_INDEX_RE = re.compile(r"\[(\d+)\]")
EXT_FILE_RE = re.compile(r"\[[^\]]*\.(xls[xmb]?|xla[m]?|xlt[xm]?|csv)\]", re.I)   # [Book.xlsx] style external path
VBA_PROC_RE = re.compile(r"^\s*(?:Public\s+|Private\s+|Friend\s+)?(?:Static\s+)?(Sub|Function|Property\s+(?:Get|Let|Set))\s+([A-Za-z_][A-Za-z0-9_]*)", re.I | re.M)
VBA_RISK_RE = re.compile(r"\b(Shell|Kill|CreateObject|GetObject|Application\.OnTime|SendKeys|Environ|FileCopy|RmDir|MkDir|Workbooks\.Open|ActiveWorkbook\.SaveAs|DisplayAlerts\s*=\s*False|On Error Resume Next|Sheets\([^)]*\)\.Delete|\.Delete)\b", re.I)

# byte patterns counted in sheet XML (trailing space excludes the plural container elements)
SHEET_XML_COUNTERS = {
    "data_validations": (b"<dataValidation ", b"<x14:dataValidation "),
    "conditional_formats": (b"<cfRule ", b"<x14:cfRule "),
    "merged_ranges": (b"<mergeCell ",),
    "hyperlinks": (b"<hyperlink ",),
    "protected": (b"<sheetProtection",),
    "freeze_panes": (b'state="frozen"',),
    "legacy_drawing": (b"<legacyDrawing",),
    "controls": (b"<control ",),
}
DIMENSION_RE = re.compile(rb'<dimension ref="([A-Z]+\d+(?::[A-Z]+\d+)?)"')


def human_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"


# ----------------------------------------------------------------------------- package-level XML

def workbook_xml(zf: zipfile.ZipFile) -> tuple[ET.Element | None, dict[str, str]]:
    try:
        wb_xml = ET.fromstring(zf.read("xl/workbook.xml"))
        wb_rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    except KeyError:
        return None, {}
    rid_to_target = {r.get("Id"): r.get("Target") for r in wb_rels.findall("rel:Relationship", NS)}
    return wb_xml, rid_to_target


def parse_sheets(wb_xml: ET.Element | None, rid_to_target: dict[str, str]) -> list[dict]:
    """Ordered sheets with name, state and part path, straight from workbook.xml."""
    out = []
    if wb_xml is None:
        return out
    sheets = wb_xml.find("m:sheets", NS)
    if sheets is None:
        return out
    for i, s in enumerate(sheets.findall("m:sheet", NS)):
        rid = s.get(f"{{{NS['r']}}}id")
        target = rid_to_target.get(rid, "") or ""
        part = target.lstrip("/")
        if not part.startswith("xl/"):
            part = "xl/" + part
        out.append({"index": i, "name": s.get("name"), "state": s.get("state", "visible"), "part": part, "sheet_id": s.get("sheetId")})
    return out


def parse_defined_names(wb_xml: ET.Element | None, sheets: list[dict]) -> list[dict]:
    out = []
    if wb_xml is None:
        return out
    dn = wb_xml.find("m:definedNames", NS)
    if dn is None:
        return out
    for d in dn.findall("m:definedName", NS):
        name = d.get("name", "")
        ref = (d.text or "").strip()
        lsid = d.get("localSheetId")
        scope = "workbook"
        if lsid is not None:
            try:
                scope = sheets[int(lsid)]["name"]
            except (ValueError, IndexError):
                scope = f"sheet#{lsid}"
        out.append({
            "name": name,
            "scope": scope,
            "refers_to": ref,
            "broken": "#REF!" in ref,
            # a workbook index [1] or a file name in brackets; plain brackets are structured references (tblPlant[Date])
            "external": bool(EXT_INDEX_RE.search(ref)) or bool(EXT_FILE_RE.search(ref)),
            "hidden": d.get("hidden") == "1",
            "builtin": name.startswith("_xlnm."),
        })
    return out


def parse_external_links(zf: zipfile.ZipFile, wb_xml: ET.Element | None, rid_to_target: dict[str, str]) -> list[dict]:
    """External link targets in workbook order (index as used in formulas: [1], [2] ...)."""
    links = []
    if wb_xml is None:
        return links
    ext_refs = wb_xml.find("m:externalReferences", NS)
    if ext_refs is None:
        return links
    for idx, er in enumerate(ext_refs.findall("m:externalReference", NS), start=1):
        rid = er.get(f"{{{NS['r']}}}id")
        part = rid_to_target.get(rid, "") or ""
        part_path = ("xl/" + part) if not part.startswith("/") else part.lstrip("/")
        rels_path = part_path.replace("externalLinks/", "externalLinks/_rels/") + ".rels"
        target, kind, sheet_names = "", "unknown", []
        try:
            rels = ET.fromstring(zf.read(rels_path))
            for r in rels.findall("rel:Relationship", NS):
                target = r.get("Target", "")
                kind = r.get("Type", "").rsplit("/", 1)[-1]
        except KeyError:
            pass
        try:
            link_xml = ET.fromstring(zf.read(part_path))
            for sd in link_xml.iter(f"{{{NS['m']}}}sheetName"):
                sheet_names.append(sd.get("val", ""))
            if link_xml.find("m:ddeLink", NS) is not None:
                kind = "ddeLink"
            if link_xml.find("m:oleLink", NS) is not None:
                kind = "oleLink"
        except KeyError:
            pass
        links.append({"index": idx, "part": part_path, "target": target, "type": kind, "sheets_referenced": sheet_names})
    return links


def resolve_link_target(target: str, base_dir: Path) -> dict:
    """Best-effort check whether the external target exists on this machine."""
    t = unquote(target)
    if t.lower().startswith("file:///"):
        t = t[8:]
    elif t.lower().startswith("file://"):
        t = "\\\\" + t[7:]
    t = t.replace("/", os.sep)
    if os.path.isabs(t) or re.match(r"^[A-Za-z]:", t) or t.startswith("\\\\"):
        cand = Path(t)
    else:
        cand = base_dir / t
    try:
        return {"resolved": cand.exists(), "path": str(cand)}
    except OSError:
        return {"resolved": False, "path": str(cand)}


def parse_connections(zf: zipfile.ZipFile) -> list[dict]:
    out = []
    if "xl/connections.xml" in zf.namelist():
        try:
            root = ET.fromstring(zf.read("xl/connections.xml"))
            for c in root.findall("m:connection", NS):
                entry = {"name": c.get("name"), "type": c.get("type"), "description": c.get("description", "")}
                db = c.find("m:dbPr", NS)
                if db is not None:
                    entry["connection"] = (db.get("connection", "") or "")[:300]
                    entry["command"] = (db.get("command", "") or "")[:300]
                out.append(entry)
        except ET.ParseError:
            pass
    return out


def parse_power_query(zf: zipfile.ZipFile) -> list[str]:
    names = []
    for n in zf.namelist():
        if n.startswith("customXml/item") and n.endswith(".xml"):
            try:
                if b"DataMashup" in zf.read(n):
                    names.append(n)
            except KeyError:
                pass
    return names


def parse_table_names(zf: zipfile.ZipFile) -> list[dict]:
    out = []
    for n in zf.namelist():
        if re.match(r"xl/tables/table\d+\.xml$", n):
            try:
                t = ET.fromstring(zf.read(n))
                out.append({"name": t.get("name") or t.get("displayName"), "ref": t.get("ref"), "part": n})
            except ET.ParseError:
                pass
    return out


def count_parts(zf: zipfile.ZipFile, prefix: str) -> int:
    return sum(1 for n in zf.namelist() if n.startswith(prefix))


def sheet_rel_counts(zf: zipfile.ZipFile) -> dict[str, dict]:
    """sheet part -> counts of charts / images / pivot tables / tables / comments via rels."""
    out: dict[str, dict] = {}
    names = set(zf.namelist())
    for n in names:
        m = re.match(r"xl/worksheets/_rels/(sheet\d+\.xml)\.rels$", n)
        if not m:
            continue
        counts = {"charts": 0, "images": 0, "pivot_tables": 0, "tables": 0, "comments": 0}
        try:
            rels = ET.fromstring(zf.read(n))
        except ET.ParseError:
            continue
        for r in rels.findall("rel:Relationship", NS):
            typ = r.get("Type", "").rsplit("/", 1)[-1]
            target = r.get("Target", "")
            if typ == "drawing":
                # relative to xl/worksheets/ (../drawings/x.xml) or absolute in the package (/xl/drawings/x.xml, as openpyxl writes)
                dpath = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/worksheets/" + target)
                drels = posixpath.join(posixpath.dirname(dpath), "_rels", posixpath.basename(dpath) + ".rels")
                if drels in names:
                    try:
                        for rr in ET.fromstring(zf.read(drels)).findall("rel:Relationship", NS):
                            t2 = rr.get("Type", "").rsplit("/", 1)[-1]
                            if t2 == "chart":
                                counts["charts"] += 1
                            elif t2 == "image":
                                counts["images"] += 1
                    except ET.ParseError:
                        pass
            elif typ == "pivotTable":
                counts["pivot_tables"] += 1
            elif typ == "table":
                counts["tables"] += 1
            elif typ == "comments":
                counts["comments"] += 1
        out["xl/worksheets/" + m.group(1)] = counts
    return out


def scan_sheet_xml(zf: zipfile.ZipFile, part: str, chunk: int = 4 * 1024 * 1024) -> dict:
    """Chunked byte scan of a sheet part: declared dimension and feature counts, bounded memory."""
    res = {k: 0 for k in SHEET_XML_COUNTERS}
    res["declared_dimension"] = ""
    res["xml_bytes"] = 0
    try:
        info = zf.getinfo(part)
    except KeyError:
        return res
    res["xml_bytes"] = info.file_size
    overlap = 64
    tail = b""
    first = True
    with zf.open(part) as fh:
        while True:
            data = fh.read(chunk)
            if not data:
                break
            buf = tail + data
            if first:
                m = DIMENSION_RE.search(buf)
                if m:
                    res["declared_dimension"] = m.group(1).decode()
                first = False
            for key, pats in SHEET_XML_COUNTERS.items():
                for p in pats:
                    res[key] += buf.count(p)
            # subtract matches that will be counted again from the overlap region
            tail = data[-overlap:] if len(data) >= overlap else data
            if len(buf) > len(data):  # remove matches fully inside the previous tail, which were counted last round
                prev_tail = buf[: len(buf) - len(data)]
                for key, pats in SHEET_XML_COUNTERS.items():
                    for p in pats:
                        res[key] -= prev_tail.count(p)
    res["protected"] = res["protected"] > 0
    res["freeze_panes"] = res["freeze_panes"] > 0
    return res


def dim_rows_cols(dim: str) -> tuple[int, int]:
    if not dim:
        return 0, 0
    last = dim.split(":")[-1]
    m = re.match(r"([A-Z]+)(\d+)", last)
    if not m:
        return 0, 0
    col = 0
    for ch in m.group(1):
        col = col * 26 + (ord(ch) - 64)
    return int(m.group(2)), col


# ----------------------------------------------------------------------------- VBA

def analyse_vba(path: Path) -> dict:
    result = {"present": False, "modules": [], "procedures": 0, "lines": 0, "auto_exec": [], "risk_hits": [], "error": None}
    if not HAVE_OLEVBA:
        result["error"] = "oletools not installed; VBA not analysed"
        return result
    try:
        vp = VBA_Parser(str(path))
    except Exception as exc:  # pragma: no cover
        result["error"] = f"olevba failed: {exc}"
        return result
    try:
        if not vp.detect_vba_macros():
            return result
        result["present"] = True
        for (_fn, _stream, vba_filename, code) in vp.extract_macros():
            code = code or ""
            lines = [ln for ln in code.splitlines() if ln.strip() and not ln.strip().startswith("Attribute ")]
            procs = [(m.group(1).split()[0], m.group(2)) for m in VBA_PROC_RE.finditer(code)]
            module = {
                "module": vba_filename,
                "lines": len(lines),
                "procedures": [f"{k} {n}" for k, n in procs],
                "sheets_referenced": sorted(set(re.findall(r"(?:Sheets|Worksheets)\(\"([^\"]+)\"\)", code))),
                "workbooks_referenced": sorted(set(re.findall(r"Workbooks(?:\.Open)?\(\"([^\"]+)\"\)", code))),
                "hardcoded_paths": sorted(set(re.findall(r"\"([A-Za-z]:\\[^\"]+|\\\\[^\"]+)\"", code))),
                "risk_keywords": sorted(set(m.group(1) for m in VBA_RISK_RE.finditer(code))),
            }
            result["modules"].append(module)
            result["procedures"] += len(procs)
            result["lines"] += len(lines)
            for _k, n in procs:
                nl = n.lower()
                if nl in ("workbook_open", "auto_open", "workbook_beforeclose", "workbook_beforesave", "auto_close") or nl.startswith("worksheet_"):
                    result["auto_exec"].append(f"{vba_filename}:{n}")
            result["risk_hits"].extend(f"{vba_filename}:{r}" for r in module["risk_keywords"])
    finally:
        vp.close()
    return result


# ----------------------------------------------------------------------------- cell scans

def scan_formula(f: str, own_sheet: str, stats: dict, deps: Counter, ext_use: Counter):
    fu = f.upper()
    stats["formulas"] += 1
    if any(v in fu for v in VOLATILE):
        stats["volatile"] += 1
    if any(l in fu for l in LOOKUPS):
        stats["lookups"] += 1
    if "#REF!" in fu:
        stats["ref_errors_in_formula"] += 1
    ext_idx = [int(m.group(1)) for m in EXT_INDEX_RE.finditer(f)]
    if ext_idx:
        stats["external_link_formulas"] += 1
        for idx in ext_idx:
            ext_use[idx] += 1
    for m in SHEET_REF_RE.finditer(f):
        name = m.group(1) or m.group(2)
        if name is None:
            continue
        name = name.replace("''", "'")
        if "[" in name and "]" in name:
            continue
        if name != own_sheet:
            deps[(own_sheet, name)] += 1


def iter_cells(ws):
    """Stream rows of a read-only sheet without padding to the declared dimension."""
    try:
        ws.reset_dimensions()
    except Exception:
        pass
    if hasattr(ws, "_cells_by_row"):
        return ws._cells_by_row(1, 1, None, None, values_only=False)
    return ws.iter_rows()


def formula_text(v) -> str | None:
    if isinstance(v, str):
        return v if v.startswith("=") else None
    txt = getattr(v, "text", None)  # ArrayFormula
    if txt is not None:
        return txt if str(txt).startswith("=") else "=" + str(txt)
    if v.__class__.__name__ in ("DataTableFormula",):
        return "=TABLE()"
    return None


def inspect_workbook(path: Path) -> dict:
    t0 = time.time()
    info: dict = {
        "file": str(path),
        "name": path.name,
        "size_bytes": path.stat().st_size,
        "size": human_size(path.stat().st_size),
        "modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M"),
        "errors": [],
    }
    if not zipfile.is_zipfile(path):
        info["errors"].append("Not an OOXML zip (legacy .xls or .xlsb are not supported by this tool; save as .xlsm)")
        return info

    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        wb_xml, rid_to_target = workbook_xml(zf)
        sheet_meta = parse_sheets(wb_xml, rid_to_target)
        info["has_vba_project"] = "xl/vbaProject.bin" in names
        info["external_links"] = parse_external_links(zf, wb_xml, rid_to_target)
        for l in info["external_links"]:
            l.update(resolve_link_target(l["target"], path.parent))
        info["defined_names"] = parse_defined_names(wb_xml, sheet_meta)
        info["connections"] = parse_connections(zf)
        info["power_query_parts"] = parse_power_query(zf)
        info["tables"] = parse_table_names(zf)
        info["part_counts"] = {
            "worksheets": count_parts(zf, "xl/worksheets/sheet"),
            "charts": count_parts(zf, "xl/charts/chart"),
            "drawings": count_parts(zf, "xl/drawings/drawing"),
            "pivot_caches": count_parts(zf, "xl/pivotCache/pivotCacheDefinition"),
            "pivot_tables": count_parts(zf, "xl/pivotTables/pivotTable"),
            "tables": count_parts(zf, "xl/tables/table"),
            "images": count_parts(zf, "xl/media/"),
            "query_tables": count_parts(zf, "xl/queryTables/"),
            "slicers": count_parts(zf, "xl/slicers/"),
            "external_link_parts": count_parts(zf, "xl/externalLinks/externalLink"),
            "activex": count_parts(zf, "xl/activeX/"),
            "embeddings": count_parts(zf, "xl/embeddings/"),
        }
        info["uncompressed_bytes"] = sum(i.file_size for i in zf.infolist())
        largest = sorted(zf.infolist(), key=lambda i: i.file_size, reverse=True)[:8]
        info["largest_parts"] = [{"part": i.filename, "size": human_size(i.file_size)} for i in largest]
        rel_counts = sheet_rel_counts(zf)
        xml_meta = {s["part"]: scan_sheet_xml(zf, s["part"]) for s in sheet_meta}

    info["vba"] = analyse_vba(path) if info["has_vba_project"] else {"present": False, "modules": [], "procedures": 0, "lines": 0, "auto_exec": [], "risk_hits": [], "error": None}

    # ---- formula pass (streaming)
    by_name = {s["name"]: s for s in sheet_meta}
    sheets: list[dict] = []
    deps: Counter = Counter()
    ext_use: Counter = Counter()
    total: Counter = Counter()
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=True)
    except Exception as exc:
        info["errors"].append(f"openpyxl could not open workbook: {exc}")
        return info
    try:
        for ws in wb.worksheets:
            meta = by_name.get(ws.title, {"state": getattr(ws, "sheet_state", "visible"), "part": ""})
            xm = xml_meta.get(meta.get("part"), {})
            st = {"formulas": 0, "volatile": 0, "lookups": 0, "ref_errors_in_formula": 0, "array_formulas": 0, "external_link_formulas": 0}
            values = 0
            last_row = 0
            last_col = 0
            samples: list[str] = []
            try:
                for row in iter_cells(ws):
                    for c in row:
                        v = getattr(c, "value", None)
                        if v is None:
                            continue
                        values += 1
                        r_, c_ = getattr(c, "row", 0), getattr(c, "column", 0)
                        if r_ > last_row:
                            last_row = r_
                        if c_ > last_col:
                            last_col = c_
                        f = formula_text(v)
                        if f is None:
                            continue
                        if not isinstance(v, str):
                            st["array_formulas"] += 1
                        scan_formula(f, ws.title, st, deps, ext_use)
                        if len(samples) < 5 and ("[" in f or "!" in f):
                            samples.append(f"{get_column_letter(c_)}{r_}: {f[:120]}")
            except Exception as exc:
                info["errors"].append(f"cell scan failed on '{ws.title}': {exc}")
            decl_rows, decl_cols = dim_rows_cols(xm.get("declared_dimension", ""))
            used = f"A1:{get_column_letter(last_col)}{last_row}" if last_row and last_col else "(empty)"
            bloated = decl_rows >= BLOAT_MIN_ROWS and decl_rows > 2 * max(last_row, 1)
            entry = {
                "sheet": ws.title,
                "state": meta.get("state", "visible"),
                "declared_dimension": xm.get("declared_dimension", ""),
                "used_range": used,
                "last_row": last_row,
                "last_col": last_col,
                "bloated_dimension": bloated,
                "xml_bytes": xm.get("xml_bytes", 0),
                "cells_with_values": values,
                **st,
                "merged_ranges": xm.get("merged_ranges", 0),
                "data_validations": xm.get("data_validations", 0),
                "conditional_formats": xm.get("conditional_formats", 0),
                "hyperlinks": xm.get("hyperlinks", 0),
                "protected": xm.get("protected", False),
                "freeze_panes": xm.get("freeze_panes", False),
                "form_controls": xm.get("controls", 0) + xm.get("legacy_drawing", 0),
                **rel_counts.get(meta.get("part", ""), {"charts": 0, "images": 0, "pivot_tables": 0, "tables": 0, "comments": 0}),
                "link_formula_samples": samples,
            }
            sheets.append(entry)
            for k, v in st.items():
                total[k] += v
            total["cells_with_values"] += values
        for cs in getattr(wb, "chartsheets", []):
            sheets.append({"sheet": cs.title, "state": by_name.get(cs.title, {}).get("state", "visible"), "type": "chartsheet"})
    finally:
        wb.close()

    info["sheets"] = sheets
    info["totals"] = dict(total)
    info["cross_sheet_dependencies"] = [{"from": a, "to": b, "formulas": n} for (a, b), n in sorted(deps.items(), key=lambda kv: -kv[1])]
    known = {s["sheet"] for s in sheets}
    info["dangling_sheet_refs"] = sorted({b for (a, b) in deps if b not in known})
    for l in info["external_links"]:
        l["formulas_using"] = ext_use.get(l["index"], 0)

    # ---- cached value error pass (streaming)
    err: Counter = Counter()
    err_by_sheet: dict[str, Counter] = defaultdict(Counter)
    try:
        wbv = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for ws in wbv.worksheets:
                for row in iter_cells(ws):
                    for c in row:
                        v = getattr(c, "value", None)
                        if isinstance(v, str) and v in ERROR_VALUES:
                            err[v] += 1
                            err_by_sheet[ws.title][v] += 1
        finally:
            wbv.close()
    except Exception as exc:
        info["errors"].append(f"cached value scan failed: {exc}")
    info["cached_errors"] = dict(err)
    info["cached_errors_by_sheet"] = {k: dict(v) for k, v in err_by_sheet.items()}
    info["scan_seconds"] = round(time.time() - t0, 1)
    return info


# ----------------------------------------------------------------------------- reporting

def md_table(headers: list[str], rows: list[list]) -> str:
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    for r in rows:
        out.append("| " + " | ".join(str(x).replace("|", "\\|") if x is not None else "" for x in r) + " |")
    return "\n".join(out)


def render_report(results: list[dict], base: Path) -> str:
    L: list[str] = []
    L.append(f"# Excel workbook inspection\n\nSource: `{base}`  \nGenerated: {datetime.now():%Y-%m-%d %H:%M}\n")

    rows = []
    for r in results:
        if "sheets" not in r:
            rows.append([r["name"], r["size"], "ERROR: " + "; ".join(r.get("errors", [])), "", "", "", "", "", ""])
            continue
        hidden = sum(1 for s in r["sheets"] if s.get("state") != "visible")
        dead = sum(1 for l in r["external_links"] if not l.get("resolved"))
        vba = r["vba"]
        rows.append([
            r["name"], r["size"], len(r["sheets"]), hidden,
            f"{len(r['external_links'])} ({dead} unresolved)",
            r["totals"].get("formulas", 0),
            r["totals"].get("external_link_formulas", 0),
            f"{vba['procedures']} procs / {vba['lines']} lines" if vba["present"] else ("bin present, unparsed" if r.get("has_vba_project") else "none"),
            sum(r["cached_errors"].values()),
        ])
    L.append("## Summary\n")
    L.append(md_table(["Workbook", "Size", "Sheets", "Hidden", "External links", "Formulas", "Ext-link formulas", "VBA", "Cached errors"], rows))
    L.append("")

    L.append("## Cross-workbook link map\n")
    names = {r["name"].lower() for r in results}
    graph_rows = []
    for r in results:
        for l in r.get("external_links", []):
            tgt = unquote(l["target"])
            tgt_name = tgt.replace("\\", "/").rsplit("/", 1)[-1]
            graph_rows.append([r["name"], f"[{l['index']}]", l.get("type", ""), tgt[:110], "ok" if l.get("resolved") else "MISSING", "yes" if tgt_name.lower() in names else "no", l.get("formulas_using", 0), ", ".join(l["sheets_referenced"][:6])])
    L.append(md_table(["From", "Idx", "Type", "Target", "On disk", "In this set", "Formulas", "Sheets referenced"], graph_rows) if graph_rows else "_No external links found._")
    L.append("")

    for r in results:
        L.append(f"\n---\n\n## {r['name']}\n")
        L.append(f"- Size {r['size']} on disk, {human_size(r.get('uncompressed_bytes', 0))} uncompressed. Modified {r['modified']}. Scan {r.get('scan_seconds', '?')} s.")
        for e in r.get("errors", []):
            L.append(f"- **Error:** {e}")
        if "sheets" not in r:
            continue
        pc = r["part_counts"]
        L.append(f"- Parts: {pc['worksheets']} worksheets, {pc['charts']} charts, {pc['pivot_tables']} pivot tables ({pc['pivot_caches']} caches), {pc['tables']} tables, {pc['images']} media files, {pc['query_tables']} query tables, {pc['slicers']} slicers, {pc['activex']} ActiveX, {pc['embeddings']} embeddings.")
        L.append(f"- Data connections: {len(r['connections'])}; Power Query parts: {len(r['power_query_parts'])}.")
        L.append("- Largest parts: " + ", ".join(f"{p['part'].split('/')[-1]} {p['size']}" for p in r["largest_parts"][:5]))
        L.append("")

        L.append("### Sheets\n")
        srows = []
        for s in r["sheets"]:
            if s.get("type") == "chartsheet":
                srows.append([s["sheet"], s["state"], "chartsheet", "", "", "", "", "", "", "", "", ""])
                continue
            used = s["used_range"] + (" BLOAT" if s["bloated_dimension"] else "")
            srows.append([
                s["sheet"], s["state"], s["declared_dimension"], used, s["cells_with_values"], s["formulas"],
                s["external_link_formulas"], s["volatile"], s["lookups"], s["data_validations"],
                s["conditional_formats"], f"{s['charts']}c/{s['pivot_tables']}p/{s['images']}i/{s['tables']}t",
            ])
        L.append(md_table(["Sheet", "State", "Declared", "Used", "Values", "Formulas", "Ext-link", "Volatile", "Lookups", "DV", "CF", "Charts/Pivots/Imgs/Tables"], srows))
        bloat = [s["sheet"] for s in r["sheets"] if s.get("bloated_dimension")]
        if bloat:
            L.append(f"\nDeclared range far larger than real content (dead rows, slows the file): {', '.join(bloat)}")
        prot = [s["sheet"] for s in r["sheets"] if s.get("protected")]
        if prot:
            L.append(f"\nProtected sheets: {', '.join(prot)}")
        L.append("")

        L.append("### Cross-sheet dependencies (formula count, top 40)\n")
        deps = r["cross_sheet_dependencies"][:40]
        L.append(md_table(["From sheet", "Reads from", "Formulas"], [[d["from"], d["to"], d["formulas"]] for d in deps]) if deps else "_None._")
        if r["dangling_sheet_refs"]:
            L.append(f"\n**Dangling sheet references (sheet no longer exists):** {', '.join(r['dangling_sheet_refs'])}")
        L.append("")

        dn = r["defined_names"]
        broken = [d for d in dn if d["broken"]]
        ext = [d for d in dn if d["external"] and not d["broken"]]
        L.append(f"### Defined names: {len(dn)} total, {len(broken)} broken (#REF!), {len(ext)} pointing to external workbooks, {sum(1 for d in dn if d['hidden'])} hidden\n")
        show = broken + ext
        if show:
            L.append(md_table(["Name", "Scope", "Refers to", "Status"], [[d["name"], d["scope"], d["refers_to"][:100], "BROKEN" if d["broken"] else "external"] for d in show[:60]]))
        L.append("")

        v = r["vba"]
        L.append("### VBA\n")
        if not v["present"]:
            if v.get("error"):
                L.append(f"_{v['error']}_")
            elif r.get("has_vba_project"):
                L.append("**vbaProject.bin is present but olevba found no parseable macros (corrupt, protected, or empty project). Open in the VBA editor to confirm.**")
            else:
                L.append("_No VBA project._")
        else:
            L.append(f"{len(v['modules'])} modules, {v['procedures']} procedures, {v['lines']} code lines.")
            if v["auto_exec"]:
                L.append(f"\n**Auto-executing procedures:** {', '.join(v['auto_exec'])}")
            if v["risk_hits"]:
                L.append(f"\n**Keywords to review:** {', '.join(sorted(set(v['risk_hits'])))}")
            L.append("")
            L.append(md_table(["Module", "Lines", "Procedures", "Sheets referenced", "Hard-coded paths"], [[m["module"], m["lines"], ", ".join(m["procedures"])[:200], ", ".join(m["sheets_referenced"])[:120], ", ".join(m["hardcoded_paths"])[:120]] for m in v["modules"]]))
        L.append("")

        if r["connections"]:
            L.append("### Data connections\n")
            L.append(md_table(["Name", "Type", "Connection", "Command"], [[c.get("name"), c.get("type"), c.get("connection", "")[:100], c.get("command", "")[:100]] for c in r["connections"]]))
            L.append("")

        if r["tables"]:
            L.append("### Excel tables (ListObjects)\n")
            L.append(", ".join(f"{t['name']} ({t['ref']})" for t in r["tables"][:40]))
            L.append("")

        if r["cached_errors"]:
            L.append("### Cached error values\n")
            L.append(md_table(["Sheet"] + list(ERROR_VALUES), [[sh] + [cnt.get(e, 0) for e in ERROR_VALUES] for sh, cnt in sorted(r["cached_errors_by_sheet"].items(), key=lambda kv: -sum(kv[1].values()))[:30]]))
            L.append("")

        samples = [(s["sheet"], f) for s in r["sheets"] for f in s.get("link_formula_samples", [])]
        if samples:
            L.append("### Sample cross-reference formulas\n")
            L.append("\n".join(f"- `{sh}` {f}" for sh, f in samples[:25]))
            L.append("")
    return "\n".join(L)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+", help="folder(s) and/or workbook file(s)")
    ap.add_argument("-o", "--out", default="xl_inspect_out", help="output directory")
    args = ap.parse_args(argv)

    files: list[Path] = []
    for p in args.paths:
        pp = Path(p)
        if pp.is_dir():
            files.extend(sorted(x for x in pp.iterdir() if x.suffix.lower() in (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".xlsb") and not x.name.startswith("~$")))
        elif pp.exists():
            files.append(pp)
        else:
            print(f"warning: {p} not found", file=sys.stderr)
    if not files:
        sys.exit("no workbooks found")

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for f in files:
        print(f"inspecting {f.name} ({human_size(f.stat().st_size)}) ...", flush=True)
        try:
            r = inspect_workbook(f)
        except Exception as exc:
            r = {"file": str(f), "name": f.name, "size": human_size(f.stat().st_size), "size_bytes": f.stat().st_size, "modified": "", "errors": [f"unhandled: {exc!r}"]}
        results.append(r)
        (out / (f.stem + ".json")).write_text(json.dumps(r, indent=2, default=str), encoding="utf-8")
        print(f"  done in {r.get('scan_seconds', '?')} s", flush=True)

    base = files[0].parent if len({f.parent for f in files}) == 1 else Path(os.path.commonpath([str(f) for f in files]))
    (out / "inspection_report.md").write_text(render_report(results, base), encoding="utf-8")
    print(f"\nreport written to {out / 'inspection_report.md'}")
    if not HAVE_OLEVBA:
        print("note: install oletools for VBA analysis:  pip install oletools")


if __name__ == "__main__":
    main()
